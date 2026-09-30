"""Text extraction and viewer payloads for supported document types."""

import csv
import html
import io
import json
from dataclasses import dataclass, field
from pathlib import Path

import mammoth
import markdown
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

KIND_BY_EXTENSION = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".xlsx": "sheet",
    ".xlsm": "sheet",
    ".csv": "sheet",
    ".txt": "text",
    ".md": "markdown",
    ".json": "text",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".webp": "image",
}

MAX_SHEET_ROWS = 500
MAX_SHEET_COLS = 40
MAX_SLIDE_IMAGES = 3


class UnsupportedDocument(ValueError):
    """Raised when a file type cannot be processed."""


@dataclass
class Section:
    label: str
    text: str


@dataclass
class ExtractedDocument:
    kind: str
    sections: list[Section] = field(default_factory=list)
    preview: dict = field(default_factory=dict)
    assets: list[tuple[str, bytes]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections if s.text)


def kind_for(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in KIND_BY_EXTENSION:
        raise UnsupportedDocument(
            f"Unsupported file type '{ext or filename}'. Supported: " + ", ".join(sorted(KIND_BY_EXTENSION))
        )
    return KIND_BY_EXTENSION[ext]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def render_markdown(text: str) -> str:
    """Render Markdown with any raw HTML escaped."""
    return markdown.markdown(html.escape(text, quote=False), extensions=["tables", "sane_lists"])


def _split_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []
    for line in text.splitlines():
        if line.lstrip().startswith("#"):
            if "".join(buffer).strip():
                sections.append(Section(heading, "\n".join(buffer).strip()))
            heading = line.lstrip("# ").strip()
            buffer = [heading]
        else:
            buffer.append(line)
    if "".join(buffer).strip():
        sections.append(Section(heading, "\n".join(buffer).strip()))
    return sections


def _extract_pdf(data: bytes) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(data))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - one malformed page must not abort the document
            text = ""
        sections.append(Section(f"Page {number}", text.strip()))
    return ExtractedDocument("pdf", sections, {"kind": "pdf", "pages": len(reader.pages)})


def _extract_docx(data: bytes) -> ExtractedDocument:
    document = DocxDocument(io.BytesIO(data))
    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.lower().startswith(("heading", "title")):
            if buffer:
                sections.append(Section(heading, "\n".join(buffer)))
            heading, buffer = text, [text]
        else:
            buffer.append(text)
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        buffer.append("\n".join(rows))
    if buffer:
        sections.append(Section(heading, "\n".join(buffer)))
    rendered = mammoth.convert_to_html(io.BytesIO(data)).value
    return ExtractedDocument("docx", sections, {"kind": "html", "html": rendered})


def _extract_pptx(data: bytes) -> ExtractedDocument:
    presentation = Presentation(io.BytesIO(data))
    sections: list[Section] = []
    slides = []
    assets: list[tuple[str, bytes]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        body: list[dict] = []
        images: list[int] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        body.append({"text": text, "level": paragraph.level})
            elif shape.has_table:
                for row in shape.table.rows:
                    text = " | ".join(cell.text.strip() for cell in row.cells)
                    if text.strip(" |"):
                        body.append({"text": text, "level": 0})
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and len(images) < MAX_SLIDE_IMAGES:
                image = shape.image
                images.append(len(assets))
                assets.append((image.content_type, image.blob))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"title": title, "body": body, "images": images, "notes": notes})
        text = "\n".join([title] + [b["text"] for b in body] + ([notes] if notes else []))
        sections.append(Section(f"Slide {number}", text.strip()))
    return ExtractedDocument("pptx", sections, {"kind": "slides", "slides": slides}, assets)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _sheet_payload(name: str, rows: list[list[str]]) -> tuple[dict, Section]:
    truncated = len(rows) > MAX_SHEET_ROWS or any(len(r) > MAX_SHEET_COLS for r in rows)
    clipped = [r[:MAX_SHEET_COLS] for r in rows[:MAX_SHEET_ROWS]]
    width = max((len(r) for r in clipped), default=0)
    clipped = [r + [""] * (width - len(r)) for r in clipped]
    text = "\n".join(" | ".join(r) for r in clipped if any(r))
    return {"name": name, "rows": clipped, "truncated": truncated}, Section(f"Sheet: {name}", text)


def _extract_sheet(data: bytes, filename: str) -> ExtractedDocument:
    sheets, sections = [], []
    if filename.lower().endswith(".csv"):
        rows = [list(r) for r in csv.reader(io.StringIO(_decode(data)))]
        sheet, section = _sheet_payload(Path(filename).stem, rows)
        sheets.append(sheet)
        sections.append(section)
    else:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                rows.append([_cell(v) for v in row])
                if len(rows) > MAX_SHEET_ROWS:
                    break
            while rows and not any(rows[-1]):
                rows.pop()
            sheet, section = _sheet_payload(worksheet.title, rows)
            sheets.append(sheet)
            sections.append(section)
        workbook.close()
    return ExtractedDocument("sheet", sections, {"kind": "sheets", "sheets": sheets})


def extract(data: bytes, filename: str) -> ExtractedDocument:
    """Extract searchable text sections and a viewer payload from a file."""
    kind = kind_for(filename)
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    if kind == "pptx":
        return _extract_pptx(data)
    if kind == "sheet":
        return _extract_sheet(data, filename)
    if kind == "image":
        return ExtractedDocument("image", [], {"kind": "image"})
    text = _decode(data)
    if kind == "markdown":
        return ExtractedDocument("markdown", _split_markdown(text), {"kind": "html", "html": render_markdown(text)})
    if filename.lower().endswith(".json"):
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except json.JSONDecodeError:
            pass
    return ExtractedDocument("text", [Section("", text.strip())], {"kind": "text", "text": text})
