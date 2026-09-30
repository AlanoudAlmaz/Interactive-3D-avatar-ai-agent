import io
import json
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from zayed.assistant import Assistant
from zayed.config import Settings
from zayed.files import FileStore
from zayed.knowledge import KnowledgeBase

APPROVAL = {
    "source": "Internal policy portal",
    "owner": "HR",
    "approved_by": "HR Director",
    "approved_on": "2026-01-01",
    "status": "approved",
}


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    doc = Document()
    for style, text in paragraphs:
        if style.startswith("h"):
            doc.add_heading(text, int(style[1:]))
        else:
            doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_pptx(slides: list[str]) -> bytes:
    prs = Presentation()
    for title in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = f"Details about {title}"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture
def knowledge_dir(tmp_path: Path) -> Path:
    root = tmp_path / "knowledge"
    (root / "approved").mkdir(parents=True)
    (root / "approved" / "leave.md").write_text(
        "# Leave policy\n\n## Annual leave\n\nEmployees receive 30 calendar days of annual leave per year.\n\n"
        "## Sick leave\n\nSick leave requires a medical certificate after two consecutive days.\n",
        encoding="utf-8",
    )
    (root / "approved" / "leave.ar.md").write_text(
        "# سياسة الإجازات\n\n## الإجازة السنوية\n\nيحصل الموظف على ثلاثين يوماً من الإجازة السنوية في كل عام.\n",
        encoding="utf-8",
    )
    (root / "approved" / "draft.md").write_text("# Draft\n\nThe cafeteria opens at 6am.\n", encoding="utf-8")
    manifest = {
        "documents": [
            {
                "id": "leave",
                "title": "Leave Policy",
                "title_ar": "سياسة الإجازات",
                "file": "approved/leave.md",
                "language": "en",
                **APPROVAL,
            },
            {
                "id": "leave-ar",
                "title": "Leave Policy (Arabic)",
                "title_ar": "سياسة الإجازات",
                "file": "approved/leave.ar.md",
                "language": "ar",
                **APPROVAL,
            },
            {"id": "draft", "title": "Draft", "file": "approved/draft.md", **APPROVAL, "status": "draft"},
            {
                "id": "incomplete",
                "title": "No approver",
                "file": "approved/draft.md",
                **{k: v for k, v in APPROVAL.items() if k != "approved_by"},
            },
            {"id": "missing", "title": "Missing", "file": "approved/nope.md", **APPROVAL},
            {"id": "escape", "title": "Escape", "file": "../outside.md", **APPROVAL},
        ]
    }
    (tmp_path / "outside.md").write_text("# Outside\n\nSecret.\n", encoding="utf-8")
    (root / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return root


@pytest.fixture
def knowledge(knowledge_dir: Path) -> KnowledgeBase:
    return KnowledgeBase(knowledge_dir).load()


@pytest.fixture
def store(tmp_path: Path) -> FileStore:
    return FileStore(tmp_path / "uploads", 1024 * 1024).load()


@pytest.fixture
def settings(knowledge_dir: Path, tmp_path: Path) -> Settings:
    return Settings(
        azure_speech_key="",
        azure_openai_endpoint="",
        azure_openai_key="",
        azure_openai_deployment="",
        openai_key="",
        knowledge_dir=knowledge_dir,
        upload_dir=tmp_path / "uploads",
    )


@pytest.fixture
def assistant(settings: Settings, knowledge: KnowledgeBase, store: FileStore) -> Assistant:
    return Assistant(settings, knowledge, store)
